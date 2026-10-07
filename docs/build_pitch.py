#!/usr/bin/env python3
# 生成「江城信约 AgentEscrow」路演 PPT，沿用姊妹项目 at_pitch / ei_pitch 的版式与配色。
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── 配色（取自参考 PPT）────────────────────────────
DARK   = RGBColor(0x18, 0x2A, 0x36)   # 标题深蓝灰
ACCENT = RGBColor(0x24, 0x5E, 0x87)   # 强调蓝
BODY   = RGBColor(0x3C, 0x4F, 0x5D)   # 正文
MUTED  = RGBColor(0x53, 0x66, 0x75)   # 次要
CARD   = RGBColor(0xEA, 0xF2, 0xF7)   # 浅蓝卡片
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
WARM   = RGBColor(0xEC, 0xEB, 0xE7)   # 暖底
RED    = RGBColor(0xA3, 0x4E, 0x39)
GREEN  = RGBColor(0x2B, 0x6D, 0x56)
FONT   = "Arial Unicode MS"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def slide():
    s = prs.slides.add_slide(BLANK)
    # 暖底
    bg = s.background
    bg.fill.solid(); bg.fill.fore_color.rgb = WARM
    # 白色内容卡
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.54), Inches(0.40), Inches(12.25), Inches(6.60))
    card.fill.solid(); card.fill.fore_color.rgb = WHITE
    card.line.fill.background()
    card.shadow.inherit = False
    return s

def txt(s, l, t, w, h, text, size, color=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    return tb

def footer(s, page, total=11):
    txt(s, 0.95, 7.08, 8.85, 0.26, f"江城信约   {page:02d} / {total:02d}", 12, DARK, bold=True)

def title(s, t, sub=None):
    txt(s, 0.96, 0.81, 11.35, 0.68, t, 30, DARK, bold=True)
    if sub:
        txt(s, 0.98, 1.57, 11.15, 0.69, sub, 17, MUTED)

def stepcard(s, l, t, num, head, body, w=2.10, h=1.54, fill=CARD):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = fill; box.line.fill.background(); box.shadow.inherit = False
    txt(s, l+0.16, t+0.14, w-0.3, 0.35, num, 15, ACCENT, bold=True)
    txt(s, l+0.16, t+0.55, w-0.3, 0.42, head, 17, DARK, bold=True)
    txt(s, l+0.16, t+0.98, w-0.3, 0.5, body, 12.5, BODY)

def arrow(s, l, t=3.48):
    ln = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(l), Inches(t), Inches(0.28), Inches(0.18))
    ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT; ln.line.fill.background(); ln.shadow.inherit = False

# ═══════════ 1 封面 ═══════════
s = slide()
txt(s, 0.96, 0.91, 11.15, 0.94, "江城信约", 48, DARK, bold=True)
txt(s, 0.99, 2.01, 10.94, 0.55, "AgentEscrow · 链上信誉 + 任务托管系统", 24, ACCENT, bold=True)
txt(s, 1.00, 3.07, 4.0, 0.45, "服务对象", 20, DARK, bold=True)
txt(s, 1.00, 3.68, 4.6, 1.04, "AI Agent 与赏金任务发布者", 18, BODY)
txt(s, 5.91, 3.07, 6.4, 0.5, "质押注册 · 托管赏金 · 验收放款 · 争议罚没", 20, DARK, bold=True)
txt(s, 5.92, 3.75, 6.3, 1.2, "EVM Layer1 · BOT Chain\n测试网 Bohr 968 / 主网 677 · 原生 BOT", 18, BODY)
txt(s, 1.00, 5.55, 11.28, 0.5, "队名：【队名】    成员：【成员】", 18, DARK, bold=True)
txt(s, 1.00, 6.10, 11.28, 0.5, "BUILD BEYOND 2026 · 从零新建赛道", 16, ACCENT, bold=True)
footer(s, 1)

# ═══════════ 2 痛点 ═══════════
s = slide(); title(s, "痛点：链上身份 ≠ 服务可信", "EIP-8004 式注册只记录「存在」，不证明「靠谱」")
pains = [
    ("刷分 / 女巫", "一次性地址刷好评，无成本、可批量复制身份"),
    ("注册≠能力", "链上声明与真实交付质量脱钩，历史评分可伪造"),
    ("信任中介贵", "跨国 Agent 交易缺最小托管，欺诈无代价"),
    ("无结算闭环", "口头交付无凭证，验收/退款/罚没无链上路径"),
]
for i,(h,b) in enumerate(pains):
    l = 1.0 + (i%2)*5.7; t = 2.6 + (i//2)*1.9
    stepcard(s, l, t, f"问题 {i+1}", h, b, w=5.2, h=1.5)
txt(s, 1.0, 6.45, 11.2, 0.5, "需要：把「信誉」和「赏金」都变成有成本、可结算、可罚没的链上资产。", 18, ACCENT, bold=True)
footer(s, 2)

# ═══════════ 3 方案总览闭环 ═══════════
s = slide(); title(s, "方案总览：信任闭环", "质押抬高作恶成本，托管把交付与放款绑定")
steps = [("01","质押注册","Agent 锁 BOT"),("02","托管赏金","发布者存 BOT"),
         ("03","接单交付","提交 hash/URI"),("04","验收放款","好评入账"),
         ("05","争议罚没","退款或 slash")]
x=1.0
for i,(n,h,b) in enumerate(steps):
    stepcard(s, x, 2.71, n, h, b)
    if i<4: arrow(s, x+2.18)
    x += 2.29
txt(s, 1.01, 4.86, 11.25, 0.5, "好评累加信誉分；差评或欺诈触发罚没，质押即「保证金」", 20, DARK, bold=True)
txt(s, 1.01, 5.61, 11.15, 0.73, "信誉分由质押量、好/差评净胜、注册时长聚合——女巫批量注册需真金白银质押。", 16, BODY)
footer(s, 3)

# ═══════════ 4 合约架构 ═══════════
s = slide(); title(s, "合约架构：两个合约", "AgentRegistry（信誉账本） + TaskEscrow（任务托管）")
stepcard(s, 1.0, 2.5, "AgentRegistry", "register(uri){质押}", "信誉账本", w=5.3, h=3.0)
txt(s, 1.2, 4.05, 5.0, 1.6, "• goodCount / badCount\n• score = 质押 + 评价 + 资历\n• slash() 罚没质押（抗女巫）\n• 信誉良好 deregister 提回", 14, BODY)
stepcard(s, 6.9, 2.5, "TaskEscrow", "createTask{value:BOT}", "任务状态机", w=5.3, h=3.0)
txt(s, 7.1, 4.05, 5.0, 1.6, "• assign / submitDeliverable\n• accept → 放款 + 好评\n• reject → Disputed 争议\n• 仲裁 resolveRelease / resolveRefund", 14, BODY)
txt(s, 1.0, 5.75, 11.2, 0.5, "Escrow 经授权调用 Registry.recordReputation；所有状态变更发 event。", 15, ACCENT, bold=True)
footer(s, 4)

# ═══════════ 5 主流程时序 ═══════════
s = slide(); title(s, "关键流程时序 · 正常路径", "发任务 → 接单 → 提交 → 验收结算")
flow = [("发布者","createTask\n托管 0.5 BOT"),("Agent","assign\n接单"),
        ("Agent","submitDeliverable\nhash+URI"),("发布者","accept\n放款+好评")]
x=1.0
for i,(who,act) in enumerate(flow):
    stepcard(s, x, 2.8, who, act.split(chr(10))[0], act.split(chr(10))[1], w=2.6, h=1.7)
    if i<3: arrow(s, x+2.68)
    x += 2.86
txt(s, 1.0, 5.1, 11.2, 0.9, "结果：Agent 收到赏金 BOT，Registry 好评 +1，信誉分上升。\n每一步链上交易均可在 scan.bohr.life / scan.botchain.ai 核验。", 16, BODY)
footer(s, 5)

# ═══════════ 6 争议罚没分支 ═══════════
s = slide(); title(s, "争议与罚没分支", "发布者拒收 → 仲裁裁决：退款 或 罚没")
stepcard(s, 1.0, 2.7, "01 reject", "进入 Disputed", "赏金冻结", w=2.6, h=1.5, fill=RGBColor(0xF7,0xE9,0xE3))
arrow(s, 3.75)
stepcard(s, 4.1, 2.7, "02 仲裁", "resolveRefund", "退款给发布者", w=2.6, h=1.5, fill=RGBColor(0xF7,0xE9,0xE3))
arrow(s, 6.85)
stepcard(s, 7.2, 2.7, "03 slash", "罚没 Agent 质押", "差评 +1", w=2.6, h=1.5, fill=RGBColor(0xF7,0xE9,0xE3))
txt(s, 1.0, 4.7, 11.2, 1.2,
    "替代路径 resolveRelease：判定交付合格 → 放款 Agent + 好评。\nMVP 治理：owner=部署者=默认仲裁人，setArbitrator 可移交。\n罚没质押进入国库，作恶成本显性化。", 16, BODY)
footer(s, 6)

# ═══════════ 7 技术要点 / 安全 ═══════════
s = slide(); title(s, "技术要点与安全", "标准 EVM / Solidity 0.8.24 / OpenZeppelin v5")
bullets = [
    ("重入保护", "ReentrancyGuard 覆盖所有放款/退款/罚没路径"),
    ("访问控制", "Ownable2Step；Escrow 唯一授权写信誉，slash 限 owner/escrow"),
    ("检查-生效-交互", "先改状态再转账，原生 BOT 低-level call 校验成功"),
    ("网络", "测试网 Bohr 968 (0x3C8) / 主网 677 (0x2A5)，RPC+env 注入"),
]
for i,(h,b) in enumerate(bullets):
    l=1.0+(i%2)*5.7; t=2.6+(i//2)*1.8
    stepcard(s, l, t, "✓", h, b, w=5.2, h=1.45, fill=RGBColor(0xE3,0xEE,0xE7))
txt(s,1.0,6.35,11.2,0.5,"前端 React19 + Vite + ethers v6，液态玻璃 UI，window.ethereum 连接钱包并自动加链。",15,ACCENT,bold=True)
footer(s, 7)

# ═══════════ 8 姊妹项目协同 ═══════════
s = slide(); title(s, "三城协同 · 故事线", "江城信约在中间：托管 → 验收 → 记账 → 复盘")
chain = [("江城信约","发任务托管\n质押 Agent",CARD),
         ("江城验真","用验真方法\n验收交付",RGBColor(0xE9,0xEE,0xF5)),
         ("江城信约","放款记信誉",CARD),
         ("江城链察","复盘链上交易",RGBColor(0xEE,0xEC,0xF0))]
x=0.9
for i,(h,b,c) in enumerate(chain):
    stepcard(s, x, 2.9, f"{i+1}", h, b, w=2.7, h=1.7, fill=c)
    if i<3: arrow(s, x+2.78)
    x+=2.92
txt(s,1.0,5.2,11.2,0.9,"验真产出的验收判定，决定信约是放款好评还是拒收争议；链察对整笔托管交易做量化复盘。",16,BODY)
footer(s, 8)

# ═══════════ 9 测试与验证 ═══════════
s = slide(); title(s, "测试与验证", "本地全流程已跑通")
txt(s,1.0,2.5,5.5,0.5,"Hardhat 测试：14 项通过",20,GREEN,bold=True)
txt(s,1.0,3.1,5.5,2.5,"• 注册 / 质押门槛 / 重复注册\n• 建任务 / 接单 / 提交\n• accept 放款 + 好评\n• reject → 争议 → 退款 + slash\n• slash 国库 / deregister 提回\n• 信誉分随好差评变化",16,BODY)
txt(s,6.8,2.5,5.5,0.5,"部署与演示",20,GREEN,bold=True)
txt(s,6.8,3.1,5.5,2.5,"• npx hardhat compile ✅\n• 本地 node 部署 + demo 脚本\n• 每笔 tx 打印 scan 链接\n• 测试网 bohr / 主网 botchain\n• 前端 typecheck + build ✅",16,BODY)
txt(s,1.0,6.0,11.2,0.6,"可核验链接样式：https://scan.bohr.life/tx/0x… 与 /address/0x…",15,ACCENT,bold=True)
footer(s, 9)

# ═══════════ 10 当前状态与后续 ═══════════
s = slide(); title(s, "当前状态与后续", "本地完成，待用户测试网账号发链")
stepcard(s,1.0,2.6,"✅ 已完成","合约 + 14 测试 + 前端","编译/测试/build 全绿",w=3.6,h=1.6)
stepcard(s,4.9,2.6,"⏳ 进行中","测试网 Bohr 968","领币→部署→demo",w=3.6,h=1.6,fill=RGBColor(0xF7,0xE9,0xE3))
stepcard(s,8.8,2.6,"➡ 后续","主网 677","Gas 申请 1 BOT 后部署",w=3.6,h=1.6)
txt(s,1.0,4.7,11.2,1.4,"待补：① 钱包连接后真实发链；② 公开 Demo URL（申请表第②栏）；\n③ 测试网合约/交易链接（第③栏）；④ 主网地址与交易记录。\nGas 申请表：https://forms.gle/7WJNKfcyJLQ4vujt7",16,BODY)
footer(s, 10)

# ═══════════ 11 尾页 ═══════════
s = slide()
txt(s, 0.96, 2.6, 11.5, 1.2, "江城信约 · 让 Agent 交易可托管、可评分、可罚没", 34, DARK, bold=True)
txt(s, 1.0, 4.0, 11.0, 0.8, "AgentEscrow on BOT Chain · EVM L1 · 原生 BOT", 20, ACCENT, bold=True)
txt(s, 1.0, 5.0, 11.0, 0.6, "谢谢聆听 · 【队名】【成员】", 18, BODY)
footer(s, 11)

out = "/home/user/Doubao/chats/38445674968970754/wuhan-agent-escrow/docs/pitch.pptx"
prs.save(out)
print("saved:", out, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
